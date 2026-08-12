from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import func
from typing import Optional
from datetime import datetime
import io, csv
from database import get_db
import models
from auth import get_current_user, require_staff

router = APIRouter(prefix="/api/reports", tags=["reports"])

PRIORITY_LABELS = {"low": "Baja", "medium": "Media", "high": "Alta", "critical": "Crítica"}


def _client_filter(q, current_user, db):
    """Restrict ticket query to what the client is allowed to see."""
    if current_user.role != models.UserRole.client:
        return q
    if current_user.company:
        ids = [r[0] for r in db.query(models.User.id).filter(
            func.lower(models.User.company) == current_user.company.lower(),
            models.User.is_active == True,
        ).all()]
        return q.filter(models.Ticket.client_id.in_(ids))
    return q.filter(models.Ticket.client_id == current_user.id)


def _ticket_row(t):
    hours = None
    if t.closed_at and t.created_at:
        hours = round((t.closed_at - t.created_at).total_seconds() / 3600, 1)

    active_visits = [v for v in t.visits if v.status != "cancelled"]
    visits_con_costo = [v for v in active_visits if v.costo is not None]
    cantidad_visitas = len(visits_con_costo)
    costo_total_visitas = float(sum(v.costo for v in visits_con_costo)) if visits_con_costo else None
    horas_visitas = round(sum(v.duration_minutes or 0 for v in active_visits) / 60, 2) if active_visits else None

    return {
        "id": t.id,
        "title": t.title,
        "description": t.description or "",
        "status": t.status_rel.name if t.status_rel else None,
        "priority": t.priority,
        "priority_label": PRIORITY_LABELS.get(t.priority, t.priority),
        "client": t.client.name if t.client else None,
        "client_company": t.client.company if t.client else None,
        "assigned_agent": t.assigned_agent.name if t.assigned_agent else None,
        "category": t.category,
        "created_at": t.created_at.isoformat() if t.created_at else None,
        "closed_at": t.closed_at.isoformat() if t.closed_at else None,
        "resolution_hours": hours,
        "cantidad_visitas": cantidad_visitas,
        "costo_total_visitas": costo_total_visitas,
        "horas_visitas": horas_visitas,
    }


@router.get("/tickets")
def report_tickets(
    date_from: Optional[datetime] = Query(None),
    date_to: Optional[datetime] = Query(None),
    status_id: Optional[int] = Query(None),
    assigned_to_id: Optional[int] = Query(None),
    client_id: Optional[int] = Query(None),
    company: Optional[str] = Query(None),
    priority: Optional[str] = Query(None),
    format: str = Query("json"),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    q = db.query(models.Ticket).filter(models.Ticket.deleted_at.is_(None))
    q = _client_filter(q, current_user, db)

    if date_from:
        q = q.filter(models.Ticket.created_at >= date_from)
    if date_to:
        q = q.filter(models.Ticket.created_at <= date_to)
    if status_id:
        q = q.filter(models.Ticket.status_id == status_id)
    if assigned_to_id and current_user.role != models.UserRole.client:
        q = q.filter(models.Ticket.assigned_to_id == assigned_to_id)
    if current_user.role != models.UserRole.client:
        if company:
            ids = [r[0] for r in db.query(models.User.id).filter(
                func.lower(models.User.company) == company.lower(),
                models.User.is_active == True,
            ).all()]
            q = q.filter(models.Ticket.client_id.in_(ids))
        elif client_id:
            q = q.filter(models.Ticket.client_id == client_id)
    if priority:
        q = q.filter(models.Ticket.priority == priority)

    tickets = q.order_by(models.Ticket.created_at.desc()).all()

    # ── CSV ──────────────────────────────────────────────────────────────────
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["ID", "Título", "Estado", "Prioridad", "Cliente",
                         "Agente Asignado", "Categoría", "Fecha Creación",
                         "Fecha Cierre", "Horas Resolución",
                         "Cantidad Visitas", "Costo Total Visitas", "Horas Visitas"])
        for t in tickets:
            r = _ticket_row(t)
            writer.writerow([
                r["id"], r["title"], r["status"], r["priority_label"],
                r["client"], r["assigned_agent"] or "Sin asignar",
                r["category"] or "",
                t.created_at.strftime("%d/%m/%Y %H:%M") if t.created_at else "",
                t.closed_at.strftime("%d/%m/%Y %H:%M") if t.closed_at else "",
                r["resolution_hours"] or "",
                r["cantidad_visitas"] or "",
                r["costo_total_visitas"] or "",
                r["horas_visitas"] or "",
            ])
        output.seek(0)
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=reporte_tickets.csv"},
        )

    # ── Excel ─────────────────────────────────────────────────────────────────
    if format == "excel":
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()

        # ── Sheet 1: Tickets ──────────────────────────────────────────────
        ws = wb.active
        ws.title = "Tickets"

        header_fill = PatternFill("solid", fgColor="1A3353")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        alt_fill    = PatternFill("solid", fgColor="F0F4F8")
        thin_border = Border(bottom=Side(style="thin", color="E5E7EB"))
        center      = Alignment(horizontal="center", vertical="center")

        cols = ["ID", "Título", "Estado", "Prioridad", "Cliente",
                "Agente Asignado", "Categoría", "Fecha Creación",
                "Fecha Cierre", "Horas Resolución",
                "Cant. Visitas", "Costo Total Visitas", "Horas Visitas"]
        for ci, h in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

        ws.row_dimensions[1].height = 24

        for ri, t in enumerate(tickets, 2):
            r = _ticket_row(t)
            row_data = [
                r["id"], r["title"], r["status"], r["priority_label"],
                r["client"], r["assigned_agent"] or "Sin asignar",
                r["category"] or "",
                t.created_at.strftime("%d/%m/%Y %H:%M") if t.created_at else "",
                t.closed_at.strftime("%d/%m/%Y %H:%M") if t.closed_at else "",
                r["resolution_hours"] or "",
                r["cantidad_visitas"] or "",
                r["costo_total_visitas"] or "",
                r["horas_visitas"] or "",
            ]
            fill = alt_fill if ri % 2 == 0 else None
            for ci, val in enumerate(row_data, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.border = thin_border
                if fill:
                    cell.fill = fill

        col_widths = [8, 40, 16, 12, 22, 22, 18, 18, 18, 14, 14, 20, 14]
        for ci, w in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(1, ci).column_letter].width = w

        ws.freeze_panes = "A2"

        # ── Sheet 2: Resumen ──────────────────────────────────────────────
        ws2 = wb.create_sheet("Resumen")
        resolved = [t for t in tickets if t.closed_at]
        avg_h = None
        if resolved:
            total_h = sum((t.closed_at - t.created_at).total_seconds() / 3600
                         for t in resolved if t.created_at)
            avg_h = round(total_h / len(resolved), 1)

        by_status   = {}
        by_priority = {}
        for t in tickets:
            sn = t.status_rel.name if t.status_rel else "Sin estado"
            by_status[sn] = by_status.get(sn, 0) + 1
            by_priority[t.priority] = by_priority.get(t.priority, 0) + 1

        def sh(row, col, val, bold=False, fill_color=None):
            c = ws2.cell(row=row, column=col, value=val)
            if bold:
                c.font = Font(bold=True)
            if fill_color:
                c.fill = PatternFill("solid", fgColor=fill_color)
                c.font = Font(bold=True, color="FFFFFF")
            return c

        sh(1, 1, "RESUMEN DEL REPORTE", bold=True, fill_color="1A3353")
        ws2.merge_cells("A1:B1")

        sh(3, 1, "Total de tickets", bold=True); sh(3, 2, len(tickets))
        sh(4, 1, "Tickets resueltos", bold=True); sh(4, 2, len(resolved))
        tasa = f"{round(len(resolved)/len(tickets)*100)}%" if tickets else "0%"
        sh(5, 1, "Tasa de resolución", bold=True); sh(5, 2, tasa)
        sh(6, 1, "Horas promedio resolución", bold=True); sh(6, 2, avg_h or "—")

        sh(8, 1, "Por Estado", bold=True, fill_color="1A3353")
        sh(8, 2, "Cantidad", bold=True, fill_color="1A3353")
        for i, (name, cnt) in enumerate(by_status.items(), 9):
            sh(i, 1, name); sh(i, 2, cnt)

        base_row = 9 + len(by_status) + 1
        sh(base_row, 1, "Por Prioridad", bold=True, fill_color="1A3353")
        sh(base_row, 2, "Cantidad", bold=True, fill_color="1A3353")
        for i, (prio, cnt) in enumerate(by_priority.items(), base_row + 1):
            sh(i, 1, PRIORITY_LABELS.get(prio, prio)); sh(i, 2, cnt)

        ws2.column_dimensions["A"].width = 30
        ws2.column_dimensions["B"].width = 14

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=reporte_tickets.xlsx"},
        )

    # ── JSON ──────────────────────────────────────────────────────────────────
    return {"tickets": [_ticket_row(t) for t in tickets], "total": len(tickets)}


@router.get("/agents")
def report_agents(
    date_from: Optional[datetime] = Query(None),
    date_to:   Optional[datetime] = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_staff),
):
    """Desempeño por agente: totales, tasa resolución, tiempo promedio, clientes atendidos."""
    q = db.query(models.Ticket).filter(models.Ticket.deleted_at.is_(None))
    if date_from:
        q = q.filter(models.Ticket.created_at >= date_from)
    if date_to:
        q = q.filter(models.Ticket.created_at <= date_to)

    tickets = q.options(
        joinedload(models.Ticket.assigned_agent),
        joinedload(models.Ticket.client),
        joinedload(models.Ticket.status_rel),
    ).all()

    # Group by agent
    agents: dict = {}
    for t in tickets:
        name = t.assigned_agent.name if t.assigned_agent else "Sin asignar"
        if name not in agents:
            agents[name] = {
                "name": name,
                "tickets": [],
                "clients": {},
                "by_priority": {"low": 0, "medium": 0, "high": 0, "critical": 0},
                "by_status": {},
                "resolution_hours": [],
            }
        a = agents[name]
        a["tickets"].append(t)

        # Client tracking
        client = t.client.name if t.client else "Sin cliente"
        company = t.client.company if t.client and t.client.company else None
        ckey = client
        if ckey not in a["clients"]:
            a["clients"][ckey] = {"name": client, "company": company, "count": 0}
        a["clients"][ckey]["count"] += 1

        # Priority / status
        prio = t.priority or "low"
        if prio in a["by_priority"]:
            a["by_priority"][prio] += 1
        sname = t.status_rel.name if t.status_rel else "Sin estado"
        a["by_status"][sname] = a["by_status"].get(sname, 0) + 1

        # Resolution time
        if t.closed_at and t.created_at:
            hours = (t.closed_at - t.created_at).total_seconds() / 3600
            a["resolution_hours"].append(hours)

    result = []
    for a in agents.values():
        total    = len(a["tickets"])
        resolved = len(a["resolution_hours"])
        avg_h    = round(sum(a["resolution_hours"]) / resolved, 1) if resolved else None
        critical_high = a["by_priority"]["critical"] + a["by_priority"]["high"]

        ticket_list = sorted([
            {
                "id":               t.id,
                "title":            t.title,
                "client":           t.client.name if t.client else None,
                "client_company":   t.client.company if t.client else None,
                "status":           t.status_rel.name if t.status_rel else None,
                "priority":         t.priority,
                "priority_label":   PRIORITY_LABELS.get(t.priority, t.priority),
                "created_at":       t.created_at.isoformat() if t.created_at else None,
                "closed_at":        t.closed_at.isoformat()  if t.closed_at  else None,
                "resolution_hours": round((t.closed_at - t.created_at).total_seconds() / 3600, 1)
                                    if t.closed_at and t.created_at else None,
            }
            for t in a["tickets"]
        ], key=lambda x: x["created_at"] or "", reverse=True)

        result.append({
            "name":            a["name"],
            "total":           total,
            "resolved":        resolved,
            "open":            total - resolved,
            "resolution_rate": round(resolved / total * 100) if total else 0,
            "avg_hours":       avg_h,
            "critical_high":   critical_high,
            "by_priority":     a["by_priority"],
            "by_status":       a["by_status"],
            "clients":         sorted(a["clients"].values(), key=lambda c: c["count"], reverse=True),
            "tickets":         ticket_list,
        })

    result.sort(key=lambda a: a["total"], reverse=True)
    return {"agents": result, "total_tickets": len(tickets)}


@router.get("/categories")
def report_categories(
    date_from: Optional[datetime] = Query(None),
    date_to:   Optional[datetime] = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(require_staff),
):
    """Tickets agrupados por categoría: totales, tasa de resolución, tiempo promedio."""
    q = db.query(models.Ticket).filter(models.Ticket.deleted_at.is_(None))
    if date_from:
        q = q.filter(models.Ticket.created_at >= date_from)
    if date_to:
        q = q.filter(models.Ticket.created_at <= date_to)

    tickets = q.options(
        joinedload(models.Ticket.assigned_agent),
        joinedload(models.Ticket.client),
        joinedload(models.Ticket.status_rel),
    ).all()

    cats: dict = {}
    for t in tickets:
        name = (t.category or "").strip() or "Sin categoría"
        if name not in cats:
            cats[name] = {
                "name": name, "tickets": [],
                "by_priority": {"low": 0, "medium": 0, "high": 0, "critical": 0},
                "resolution_hours": [],
            }
        c = cats[name]
        c["tickets"].append(t)
        prio = t.priority or "low"
        if prio in c["by_priority"]:
            c["by_priority"][prio] += 1
        if t.closed_at and t.created_at:
            c["resolution_hours"].append((t.closed_at - t.created_at).total_seconds() / 3600)

    result = []
    for c in cats.values():
        total    = len(c["tickets"])
        resolved = len(c["resolution_hours"])
        avg_h    = round(sum(c["resolution_hours"]) / resolved, 1) if resolved else None
        ticket_list = sorted([
            {
                "id":               t.id,
                "title":            t.title,
                "client":           t.client.name if t.client else None,
                "client_company":   t.client.company if t.client else None,
                "agent":            t.assigned_agent.name if t.assigned_agent else None,
                "status":           t.status_rel.name if t.status_rel else None,
                "priority":         t.priority,
                "priority_label":   PRIORITY_LABELS.get(t.priority, t.priority),
                "created_at":       t.created_at.isoformat() if t.created_at else None,
                "closed_at":        t.closed_at.isoformat()  if t.closed_at  else None,
                "resolution_hours": round((t.closed_at - t.created_at).total_seconds() / 3600, 1)
                                    if t.closed_at and t.created_at else None,
            }
            for t in c["tickets"]
        ], key=lambda x: x["created_at"] or "", reverse=True)
        result.append({
            "name":            c["name"],
            "total":           total,
            "resolved":        resolved,
            "open":            total - resolved,
            "resolution_rate": round(resolved / total * 100) if total else 0,
            "avg_hours":       avg_h,
            "critical_high":   c["by_priority"]["critical"] + c["by_priority"]["high"],
            "by_priority":     c["by_priority"],
            "tickets":         ticket_list,
        })

    result.sort(key=lambda c: c["total"], reverse=True)
    return {"categories": result, "total_tickets": len(tickets)}


@router.get("/summary")
def report_summary(
    date_from: Optional[datetime] = Query(None),
    date_to: Optional[datetime] = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    q = db.query(models.Ticket).filter(models.Ticket.deleted_at.is_(None))
    q = _client_filter(q, current_user, db)
    if date_from:
        q = q.filter(models.Ticket.created_at >= date_from)
    if date_to:
        q = q.filter(models.Ticket.created_at <= date_to)

    tickets = q.all()
    by_status   = {}
    by_priority = {}
    by_agent    = {}

    for t in tickets:
        sn = t.status_rel.name if t.status_rel else "Sin estado"
        by_status[sn] = by_status.get(sn, 0) + 1
        by_priority[t.priority] = by_priority.get(t.priority, 0) + 1
        an = t.assigned_agent.name if t.assigned_agent else "Sin asignar"
        by_agent[an] = by_agent.get(an, 0) + 1

    resolved = [t for t in tickets if t.closed_at]
    avg_hours = None
    if resolved:
        total_h = sum((t.closed_at - t.created_at).total_seconds() / 3600
                      for t in resolved if t.created_at)
        avg_hours = round(total_h / len(resolved), 1)

    return {
        "total": len(tickets),
        "resolved": len(resolved),
        "avg_resolution_hours": avg_hours,
        "by_status": by_status,
        "by_priority": by_priority,
        "by_agent": by_agent,
    }
